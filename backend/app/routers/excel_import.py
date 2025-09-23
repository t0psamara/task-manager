"""
Роутер для импорта данных из Excel файлов
"""
import os
import tempfile
from typing import Dict, Any
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from sqlalchemy.ext.asyncio import AsyncSession

from ..database import get_db
from ..services.excel_import import ExcelImportService

router = APIRouter(prefix="/excel-import", tags=["excel-import"])


@router.post("/upload")
async def upload_excel_file(
    board_id: int = Form(...),
    sprint_id: int = Form(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db)
) -> Dict[str, Any]:
    """
    Загрузка и обработка Excel файла для создания фич и задач
    
    Args:
        board_id: ID доски для создания фич
        sprint_id: ID спринта для добавления задач
        file: Excel файл (.xlsx)
        db: Сессия базы данных
        
    Returns:
        Результат импорта с информацией о созданных фичах и задачах
    """
    try:
        # Проверяем тип файла
        if not file.filename.endswith('.xlsx'):
            raise HTTPException(
                status_code=400, 
                detail="Поддерживаются только файлы формата .xlsx"
            )
        
        # Проверяем размер файла (максимум 10MB)
        max_file_size = 10 * 1024 * 1024  # 10MB
        file_content = await file.read()
        
        if len(file_content) > max_file_size:
            raise HTTPException(
                status_code=400,
                detail="Размер файла не должен превышать 10MB"
            )
        
        # Сохраняем файл во временную директорию
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as temp_file:
            temp_file.write(file_content)
            temp_file_path = temp_file.name
        
        try:
            # Создаем сервис для импорта
            import_service = ExcelImportService(db)
            
            # Выполняем импорт
            result = await import_service.import_from_excel(temp_file_path, board_id, sprint_id)
            
            return {
                "message": "Импорт завершен успешно",
                "filename": file.filename,
                "board_id": board_id,
                "sprint_id": sprint_id,
                **result
            }
            
        finally:
            # Удаляем временный файл
            try:
                os.unlink(temp_file_path)
            except OSError:
                pass  # Игнорируем ошибки удаления временного файла
                
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Ошибка при обработке файла: {str(e)}"
        )


@router.post("/validate")
async def validate_excel_file(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db)
) -> Dict[str, Any]:
    """
    Валидация Excel файла без создания данных
    
    Args:
        file: Excel файл (.xlsx)
        db: Сессия базы данных
        
    Returns:
        Информация о структуре файла и потенциальных проблемах
    """
    try:
        from openpyxl import load_workbook
        
        # Проверяем тип файла
        if not file.filename.endswith('.xlsx'):
            raise HTTPException(
                status_code=400, 
                detail="Поддерживаются только файлы формата .xlsx"
            )
        
        # Сохраняем файл во временную директорию
        file_content = await file.read()
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as temp_file:
            temp_file.write(file_content)
            temp_file_path = temp_file.name
        
        try:
            # Загружаем и анализируем файл
            workbook = load_workbook(temp_file_path, read_only=True)
            
            analysis = {
                "filename": file.filename,
                "sheets_count": len(workbook.sheetnames),
                "sheets": [],
                "warnings": [],
                "errors": []
            }
            
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                sheet_info = {
                    "name": sheet_name,
                    "feature_name": None,
                    "tasks_count": 0,
                    "has_bugfix": False,
                    "warnings": []
                }
                
                # Проверяем название фичи в B1
                feature_name_cell = sheet['B1']
                if feature_name_cell.value:
                    sheet_info["feature_name"] = str(feature_name_cell.value).strip()
                else:
                    sheet_info["warnings"].append("Пустое название фичи в ячейке B1")
                
                # Проверяем данные для Багофикс в строке 3
                bugfix_cells = [sheet.cell(row=3, column=col).value for col in [3, 4, 5]]
                if any(cell is not None for cell in bugfix_cells):
                    sheet_info["has_bugfix"] = True
                
                # Подсчитываем потенциальные задачи (строки с 5-й)
                row_num = 5
                while row_num <= sheet.max_row:
                    # Проверяем есть ли данные в колонке B (название)
                    name_cell = sheet.cell(row=row_num, column=2)
                    if name_cell.value and str(name_cell.value).strip():
                        # Проверяем колонку A - если не пустая, это будет пропущено
                        skip_cell = sheet.cell(row=row_num, column=1)
                        if not (skip_cell.value and str(skip_cell.value).strip()):
                            sheet_info["tasks_count"] += 1
                    row_num += 1
                
                analysis["sheets"].append(sheet_info)
            
            return analysis
            
        finally:
            # Удаляем временный файл
            try:
                os.unlink(temp_file_path)
            except OSError:
                pass
                
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=500,
            detail=f"Ошибка при анализе файла: {str(e)}"
        )
