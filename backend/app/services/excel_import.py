"""
Сервис для импорта фич и задач из Excel файлов
"""
from typing import List, Dict, Any, Optional, Tuple
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from .. import models


class ExcelImportService:
    """Сервис для импорта данных из Excel файлов Synology"""
    
    def __init__(self, db: AsyncSession):
        self.db = db
    
    async def import_from_excel(self, file_path: str, board_id: int, sprint_id: Optional[int] = None) -> Dict[str, Any]:
        """
        Импорт фич и задач из Excel файла
        
        Args:
            file_path: Путь к Excel файлу
            board_id: ID доски для создания фич
            sprint_id: ID спринта для добавления задач (опционально)
            
        Returns:
            Результат импорта с созданными фичами и задачами
        """
        try:
            # Проверяем существование доски
            board_query = await self.db.execute(
                select(models.Board).where(models.Board.id == board_id)
            )
            board = board_query.scalar_one_or_none()
            if not board:
                raise ValueError(f"Доска с ID {board_id} не найдена")
            
            # Проверяем существование спринта если указан
            if sprint_id:
                sprint_query = await self.db.execute(
                    select(models.Sprint).where(
                        models.Sprint.id == sprint_id,
                        models.Sprint.board_id == board_id
                    )
                )
                sprint = sprint_query.scalar_one_or_none()
                if not sprint:
                    raise ValueError(f"Спринт с ID {sprint_id} не найден на доске {board_id}")
            
            # Загружаем Excel файл
            workbook = load_workbook(file_path, read_only=True)
            
            created_features = []
            created_tasks = []
            errors = []
            
            # Обрабатываем каждую вкладку как отдельную фичу
            for sheet_name in workbook.sheetnames:
                try:
                    sheet = workbook[sheet_name]
                    feature_data = await self._parse_sheet(sheet, board_id)
                    
                    if feature_data:
                        # Создаем фичу
                        feature = await self._create_feature(feature_data["feature"], board_id)
                        created_features.append(feature)
                        
                        # Создаем задачи для фичи
                        for task_data in feature_data["tasks"]:
                            task = await self._create_task(task_data, feature.id, sprint_id)
                            created_tasks.append(task)
                            
                except Exception as e:
                    error_msg = f"Ошибка при обработке вкладки '{sheet_name}': {str(e)}"
                    errors.append(error_msg)
                    print(f"[ERROR] {error_msg}")
            
            await self.db.commit()
            
            return {
                "success": True,
                "features_created": len(created_features),
                "tasks_created": len(created_tasks),
                "errors": errors,
                "details": {
                    "features": [{"id": f.id, "name": f.name} for f in created_features],
                    "tasks": [{"id": t.id, "name": t.name, "feature_id": t.feature_id} for t in created_tasks]
                }
            }
            
        except Exception as e:
            await self.db.rollback()
            raise Exception(f"Ошибка при импорте Excel файла: {str(e)}")
    
    async def _parse_sheet(self, sheet: Worksheet, board_id: int) -> Optional[Dict[str, Any]]:
        """
        Парсинг одной вкладки Excel
        
        Args:
            sheet: Вкладка Excel
            board_id: ID доски
            
        Returns:
            Данные фичи и задач или None если вкладка пустая
        """
        # Получаем название фичи из ячейки B1
        feature_name_cell = sheet['B1']
        feature_name = feature_name_cell.value
        
        if not feature_name or str(feature_name).strip() == "":
            print(f"[WARNING] Вкладка '{sheet.title}' пропущена: пустое название фичи в B1")
            return None
        
        feature_name = str(feature_name).strip()
        
        # Данные фичи
        feature_data = {
            "name": feature_name,
            "order": 0,  # Можно добавить логику для определения порядка
        }
        
        tasks_data = []
        
        # Получаем данные для "Багофикс" тикета из строки 3
        bugfix_task = self._parse_bugfix_task(sheet)
        if bugfix_task:
            tasks_data.append(bugfix_task)
        
        # Парсим тикеты начиная с 5-й строки
        row_num = 5
        while True:
            # Проверяем есть ли данные в строке
            if not self._has_data_in_row(sheet, row_num):
                break
                
            # Проверяем колонку A - если не пустая, пропускаем строку
            cell_a = sheet.cell(row=row_num, column=1)  # Колонка A
            if cell_a.value is not None and str(cell_a.value).strip() != "":
                row_num += 1
                continue
            
            # Парсим данные тикета
            task_data = self._parse_task_row(sheet, row_num)
            if task_data:
                tasks_data.append(task_data)
            
            row_num += 1
        
        return {
            "feature": feature_data,
            "tasks": tasks_data
        }
    
    def _parse_bugfix_task(self, sheet: Worksheet) -> Optional[Dict[str, Any]]:
        """
        Создание тикета "Багофикс" с данными из строки 3
        
        Args:
            sheet: Вкладка Excel
            
        Returns:
            Данные тикета "Багофикс" или None
        """
        try:
            # Загрузка для "Багофикс" из строки 3
            ios_load = self._get_cell_float_value(sheet, 3, 3)  # C3
            android_load = self._get_cell_float_value(sheet, 3, 4)  # D3
            qa_load = self._get_cell_float_value(sheet, 3, 5)  # E3
            analytics_load = 0.0  # По требованию всегда 0
            
            return {
                "name": "Багофикс",
                "estimate_ios": ios_load,
                "estimate_android": android_load,
                "estimate_qa": qa_load,
                "estimate_sa": analytics_load,
                "color": "#ffeb3b",  # Цвет по умолчанию
                "position_x": 0.0,
                "position_y": 0.0
            }
            
        except Exception as e:
            print(f"[WARNING] Ошибка при создании тикета 'Багофикс': {str(e)}")
            import traceback
            print(f"[DEBUG] Полная трассировка: {traceback.format_exc()}")
            return None
    
    def _parse_task_row(self, sheet: Worksheet, row_num: int) -> Optional[Dict[str, Any]]:
        """
        Парсинг одной строки для создания тикета
        
        Args:
            sheet: Вкладка Excel
            row_num: Номер строки
            
        Returns:
            Данные тикета или None если строка пустая
        """
        try:
            # Название из колонки B
            name_cell = sheet.cell(row=row_num, column=2)  # Колонка B
            name = name_cell.value
            
            if not name or str(name).strip() == "":
                return None
                
            name = str(name).strip()
            
            # Загрузки из колонок C, D, E, F
            ios_load = self._get_cell_float_value(sheet, row_num, 3)  # C
            android_load = self._get_cell_float_value(sheet, row_num, 4)  # D
            qa_load = self._get_cell_float_value(sheet, row_num, 5)  # E
            analytics_load = self._get_cell_float_value(sheet, row_num, 6)  # F
            
            return {
                "name": name,
                "estimate_ios": ios_load,
                "estimate_android": android_load,
                "estimate_qa": qa_load,
                "estimate_sa": analytics_load,
                "color": "#ffeb3b",  # Цвет по умолчанию
                "position_x": 0.0,
                "position_y": 0.0
            }
            
        except Exception as e:
            print(f"[WARNING] Ошибка при парсинге строки {row_num}: {str(e)}")
            return None
    
    def _get_cell_float_value(self, sheet: Worksheet, row: int, col: int) -> float:
        """
        Получение числового значения из ячейки
        
        Args:
            sheet: Вкладка Excel
            row: Номер строки
            col: Номер колонки
            
        Returns:
            Числовое значение или 0.0 если ячейка пустая/невалидная
        """
        try:
            cell = sheet.cell(row=row, column=col)
            value = cell.value
            
            if value is None:
                return 0.0
                
            if isinstance(value, (int, float)):
                # Округляем до 2 знаков после запятой для избежания проблем с точностью
                return round(float(value), 2)
                
            # Пытаемся преобразовать строку в число
            str_value = str(value).strip()
            if str_value == "":
                return 0.0
                
            return round(float(str_value), 2)
            
        except (ValueError, TypeError):
            return 0.0
    
    def _has_data_in_row(self, sheet: Worksheet, row_num: int, max_col: int = 10) -> bool:
        """
        Проверка есть ли данные в строке
        
        Args:
            sheet: Вкладка Excel
            row_num: Номер строки
            max_col: Максимальная колонка для проверки
            
        Returns:
            True если в строке есть данные
        """
        try:
            for col in range(1, max_col + 1):
                cell = sheet.cell(row=row_num, column=col)
                if cell.value is not None and str(cell.value).strip() != "":
                    return True
            return False
        except:
            return False
    
    async def _create_feature(self, feature_data: Dict[str, Any], board_id: int) -> models.Feature:
        """
        Создание фичи в базе данных
        
        Args:
            feature_data: Данные фичи
            board_id: ID доски
            
        Returns:
            Созданная фича
        """
        # Определяем порядок для новой фичи
        order_query = await self.db.execute(
            select(models.Feature.order)
            .where(models.Feature.board_id == board_id)
            .order_by(models.Feature.order.desc())
            .limit(1)
        )
        last_order = order_query.scalar_one_or_none()
        new_order = (last_order or 0) + 1
        
        feature = models.Feature(
            board_id=board_id,
            name=feature_data["name"],
            order=new_order,
            mgmt_link="",
            mgmt_title="МГМТ",
            epic_link="",
            epic_title="Эпик",
            project_code=""
        )
        
        self.db.add(feature)
        await self.db.flush()
        await self.db.refresh(feature)
        
        return feature
    
    async def _create_task(self, task_data: Dict[str, Any], feature_id: int, sprint_id: Optional[int] = None) -> models.Task:
        """
        Создание задачи в базе данных
        
        Args:
            task_data: Данные задачи
            feature_id: ID фичи
            sprint_id: ID спринта (опционально)
            
        Returns:
            Созданная задача
        """
        
        
        task = models.Task(
            feature_id=feature_id,
            sprint_id=sprint_id,  # Задачи добавляются в указанный спринт
            name=task_data["name"],
            estimate_ios=task_data["estimate_ios"],
            estimate_android=task_data["estimate_android"],
            estimate_qa=task_data["estimate_qa"],
            estimate_sa=task_data["estimate_sa"],
            position_x=task_data["position_x"],
            position_y=task_data["position_y"],
            color=task_data["color"],
            enabler_title="",
            enabler_active=True,
            is_collapsed_feature=False,
            original_feature_tasks=None
        )
        
        self.db.add(task)
        await self.db.flush()
        await self.db.refresh(task)
        
        return task
